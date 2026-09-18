"""Xuất kết quả khảo sát ra Excel (2 sheet: danh sách lượt khảo sát + chi tiết câu trả lời)."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="DDE7F3")


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL


def _format_answer(answer) -> str:
    if answer.option_id and answer.option:
        return answer.option.option_text
    if answer.answer_text is not None:
        if answer.answer_text == "yes":
            return "Có"
        if answer.answer_text == "no":
            return "Không"
        return answer.answer_text
    if answer.answer_number is not None:
        return str(answer.answer_number)
    return ""


def _answer_score(answer):
    if answer.option_id and answer.option and answer.option.score is not None:
        return answer.option.score
    if answer.question and answer.question.question_type == "rating":
        return answer.answer_number
    return ""


def build_workbook(survey, responses: list) -> io.BytesIO:
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Danh sách lượt khảo sát"
    headers1 = [
        "STT", "Ngày khảo sát", "Chi nhánh", "Dịch vụ", "Quầy", "Cán bộ",
        "Người trả lời", "Số điện thoại", "Email", "Địa chỉ",
    ]
    ws1.append(headers1)
    for i, r in enumerate(responses, start=1):
        ws1.append(
            [
                i,
                r.submitted_at.strftime("%d/%m/%Y %H:%M") if r.submitted_at else "",
                r.branch.name if r.branch else "",
                r.service_id or "",
                r.counter_id or "",
                r.employee.full_name if r.employee else "",
                r.respondent_name or "",
                r.respondent_phone or "",
                r.respondent_email or "",
                r.respondent_address or "",
            ]
        )
    _style_header(ws1)
    for col in range(1, len(headers1) + 1):
        ws1.column_dimensions[get_column_letter(col)].width = 20

    ws2 = wb.create_sheet("Chi tiết câu trả lời")
    headers2 = [
        "STT", "Ngày khảo sát", "Chi nhánh", "Cán bộ", "Khảo sát", "Câu hỏi", "Trả lời", "Điểm",
    ]
    ws2.append(headers2)
    stt = 0
    for r in responses:
        for a in r.answers:
            stt += 1
            ws2.append(
                [
                    stt,
                    r.submitted_at.strftime("%d/%m/%Y %H:%M") if r.submitted_at else "",
                    r.branch.name if r.branch else "",
                    r.employee.full_name if r.employee else "",
                    survey.title,
                    a.question.question_text if a.question else "",
                    _format_answer(a),
                    _answer_score(a),
                ]
            )
    _style_header(ws2)
    for col in range(1, len(headers2) + 1):
        ws2.column_dimensions[get_column_letter(col)].width = 32

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_summary_workbook(survey, stats: dict, *, scope_label: str | None = None) -> io.BytesIO:
    """Xuất bảng TỔNG HỢP tỷ lệ theo từng câu hỏi (không có thông tin người trả lời).

    `stats` là kết quả của `survey_statistics_service.get_statistics()` — tôn
    trọng đúng bộ lọc (thời gian/chi nhánh) đã áp dụng khi gọi, nên khi lọc theo
    1 chi nhánh thì sheet "Tổng hợp" chỉ phản ánh chi nhánh đó; không lọc thì là
    số liệu toàn hệ thống.
    """
    wb = Workbook()
    overview = stats["overview"]

    ws1 = wb.active
    ws1.title = "Tổng hợp"
    ws1.append([survey.title])
    ws1.append([f"Phạm vi: {scope_label or 'Toàn hệ thống'}"])
    ws1.append([f"Tổng số lượt khảo sát: {overview['total_responses']}"])
    avg = overview.get("average_score")
    ws1.append(
        [
            f"Điểm trung bình: {avg if avg is not None else '—'}/5"
            f" · Tỷ lệ hài lòng: {overview['satisfaction_rate']}%"
            f" · Tỷ lệ không hài lòng: {overview['dissatisfaction_rate']}%"
        ]
    )
    ws1.append([])
    header_row = ws1.max_row + 1
    headers = ["Phần", "Câu hỏi", "Phương án / Chỉ số", "Điểm", "Số lượng", "Tỷ lệ (%)"]
    ws1.append(headers)
    for cell in ws1[header_row]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL

    for q in stats["by_question"]:
        section = q.get("section") or ""
        qtext = q["question_text"]
        st = q["stats"]
        first = True

        def _row(option_label, score, count, percentage):
            nonlocal first
            ws1.append(
                [
                    section if first else "",
                    qtext if first else "",
                    option_label,
                    score if score != "" else "",
                    count if count != "" else "",
                    percentage if percentage != "" else "",
                ]
            )
            first = False

        if st["type"] == "choice":
            if st["options"]:
                for opt in st["options"]:
                    _row(
                        opt["option_text"],
                        opt.get("score") if opt.get("score") is not None else "",
                        opt["count"],
                        opt["percentage"],
                    )
                if st.get("average_score") is not None:
                    _row("Điểm trung bình", st["average_score"], "", "")
            else:
                _row("(chưa có phản hồi)", "", 0, 0)
        elif st["type"] == "yes_no":
            _row("Có", "", st["yes_count"], st["yes_percentage"])
            _row("Không", "", st["no_count"], st["no_percentage"])
        elif st["type"] == "rating":
            for b in st["breakdown"]:
                _row(b["label"], b["level"], b["count"], b["percentage"])
            _row("Điểm trung bình", st["average"] if st["average"] is not None else "", "", "")
        elif st["type"] == "number":
            _row("Trung bình", st["average"] if st["average"] is not None else "", "", "")
            _row("Nhỏ nhất", st["min"] if st["min"] is not None else "", "", "")
            _row("Lớn nhất", st["max"] if st["max"] is not None else "", "", "")
        else:  # open_text
            _row(f"{st['total_respondents']} ý kiến (xem chi tiết trong hệ thống)", "", "", "")

    for col, width in zip(range(1, 7), (22, 46, 40, 12, 12, 12)):
        ws1.column_dimensions[get_column_letter(col)].width = width

    ws2 = wb.create_sheet("Theo chi nhánh")
    headers2 = [
        "Xếp hạng", "Chi nhánh", "Tổng lượt", "Số đáp án có điểm", "Điểm trung bình",
        "Tỷ lệ hài lòng (%)", "Tỷ lệ không hài lòng (%)",
    ]
    ws2.append(headers2)
    for b in stats["by_branch"]:
        ws2.append(
            [
                b.get("rank") or "",
                b["branch_name"],
                b["total_responses"],
                b.get("scored_answers", 0),
                b["average_score"] if b["average_score"] is not None else "",
                b["satisfaction_rate"],
                b["dissatisfaction_rate"],
            ]
        )
    _style_header(ws2)
    for col in range(1, len(headers2) + 1):
        ws2.column_dimensions[get_column_letter(col)].width = 24

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
